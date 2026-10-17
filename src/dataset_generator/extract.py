"""Извлечение текста из документов: статей, книг, отчётов, веб-страниц и т.д."""

import re
import shutil
import subprocess
import zipfile
from collections import Counter
from pathlib import Path

import pymupdf
from charset_normalizer import from_bytes
from docx import Document

# Форматы, которые открывает MuPDF.
MUPDF_SUFFIXES = {
    ".pdf",
    ".epub",
    ".mobi",
    ".fb2",
    ".xps",
    ".oxps",
    ".cbz",
    ".html",
    ".htm",
    ".xhtml",
}
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".text"}
SUPPORTED_SUFFIXES = MUPDF_SUFFIXES | TEXT_SUFFIXES | {".docx", ".djvu", ".djv", ".zip"}

_BLOCK_TEXT_FLAGS = pymupdf.TEXTFLAGS_BLOCKS | pymupdf.TEXT_DEHYPHENATE
# Колонтитул: до 150 символов в первых/последних двух блоках страницы, минимум на 3 страницах.
_EDGE_BLOCKS = 2
_RUNNING_HEAD_CHARS = 150
_RUNNING_HEAD_PAGES = 3
_DIGITS = re.compile(r"\d+")
_PAGE_NUMBER = re.compile(r"^\s*[-–—]?\s*\d{1,4}\s*[-–—]?\s*$")
_HYPHEN_BREAK = re.compile(r"(\w)[-­]\n(\w)")
_SPACES = re.compile(r"[ \t ]+")


class ExtractionError(Exception):
    pass


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_SUFFIXES


def extract_text(path: Path) -> str:
    """Возвращает текст документа: один абзац (или строка исходника) на строку."""
    suffix = path.suffix.lower()
    if suffix in MUPDF_SUFFIXES:
        text = _extract_mupdf(path=path)
    elif suffix in TEXT_SUFFIXES:
        text = _decode(path.read_bytes())
    elif suffix == ".docx":
        text = _extract_docx(path)
    elif suffix in {".djvu", ".djv"}:
        text = _extract_djvu(path)
    elif suffix == ".zip":
        text = _extract_zip(path)
    else:
        raise ExtractionError(f'unsupported format "{suffix}"')

    text = clean_text(text)
    if not text:
        raise ExtractionError("the file has no text layer (probably a scan — run OCR on it first)")
    return text


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n")
    text = _HYPHEN_BREAK.sub(r"\1\2", text).replace("­", "")
    lines = []
    for line in text.split("\n"):
        line = _SPACES.sub(" ", line).strip()
        if line and not _PAGE_NUMBER.match(line):
            lines.append(line)
    return "\n".join(lines)


def _extract_mupdf(
    *, path: Path | None = None, data: bytes | None = None, filetype: str = ""
) -> str:
    try:
        doc = pymupdf.open(path) if path else pymupdf.open(stream=data, filetype=filetype)
    except Exception as exc:
        raise ExtractionError(f"cannot open the document: {exc}") from exc
    with doc:
        if doc.needs_pass:
            raise ExtractionError("the document is password-protected")
        pages = [
            [
                text
                # block = (x0, y0, x1, y1, text, block_no, block_type); type 1 — картинка
                for block in page.get_text("blocks", flags=_BLOCK_TEXT_FLAGS)
                if block[6] == 0 and (text := " ".join(block[4].split()))
            ]
            for page in doc
        ]
        # Колонтитулы есть только у документов с фиксированной вёрсткой: у EPUB/FB2/HTML
        # страницы нарезает сам MuPDF, и край страницы — случайный абзац.
        if not doc.is_reflowable:
            pages = _strip_running_heads(pages)
        return "\n".join(text for blocks in pages for text in blocks)


def _strip_running_heads(pages: list[list[str]]) -> list[list[str]]:
    """Убирает колонтитулы: короткие блоки у края страницы, повторяющиеся на разных страницах.

    В статьях это название журнала, DOI, фамилии авторов, номера страниц; в книгах —
    название книги или главы. Цифры не учитываем, чтобы «С. 12» и «С. 13» совпадали.
    """

    def key(text: str) -> str:
        return _DIGITS.sub("#", text.lower())

    def edges(blocks: list[str]) -> set[int]:
        count = len(blocks)
        return {i for i in range(count) if i < _EDGE_BLOCKS or i >= count - _EDGE_BLOCKS}

    seen: Counter[str] = Counter()
    for blocks in pages:
        seen.update(
            {key(blocks[i]) for i in edges(blocks) if len(blocks[i]) <= _RUNNING_HEAD_CHARS}
        )
    repeated = {text for text, pages_count in seen.items() if pages_count >= _RUNNING_HEAD_PAGES}
    return [
        [
            text
            for i, text in enumerate(blocks)
            if not (i in edges(blocks) and key(text) in repeated)
        ]
        for blocks in pages
    ]


def _decode(data: bytes) -> str:
    # Русские txt-файлы часто в cp1251/koi8-r — определяем кодировку автоматически.
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        match = from_bytes(data).best()
        return str(match) if match else data.decode("utf-8", errors="replace")


def _extract_docx(path: Path) -> str:
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _extract_djvu(path: Path) -> str:
    djvutxt = shutil.which("djvutxt")
    if not djvutxt:
        raise ExtractionError("DjVu needs the djvutxt tool: sudo dnf install djvulibre")
    result = subprocess.run([djvutxt, str(path)], capture_output=True, check=False)
    if result.returncode != 0:
        raise ExtractionError(f"djvutxt: {result.stderr.decode(errors='replace').strip()}")
    return _decode(result.stdout)


def _extract_zip(path: Path) -> str:
    # Типичный случай — книги вида book.fb2.zip.
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            suffix = Path(name).suffix.lower()
            if suffix in MUPDF_SUFFIXES:
                return _extract_mupdf(data=archive.read(name), filetype=suffix.lstrip("."))
            if suffix in TEXT_SUFFIXES:
                return _decode(archive.read(name))
    raise ExtractionError("the archive contains no supported files")
